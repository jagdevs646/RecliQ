"""Run the two remaining Phase 2 acceptance checks in an isolated local setup.

    backend/.venv/Scripts/python scripts/perf/phase2_acceptance.py

The API check creates a temporary 500k-row Parquet report and a temporary SQLite
database, then measures the real summary and preview routes through FastAPI's
ASGI test client. The report check runs one generated keyed benchmark through
the current and pre-Phase-2 report writers and compares workbook cell values.
Artifacts for manual visual review are written under scripts/perf/out/.
"""
from __future__ import annotations

import argparse
import gc
import importlib.util
import json
import math
import os
import re
import statistics
import subprocess
import sys
import tempfile
import threading
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

REPO = Path(__file__).resolve().parents[2]
BACKEND = REPO / "backend"
PERF = REPO / "scripts" / "perf"
sys.path.insert(0, str(BACKEND))
sys.path.insert(0, str(PERF))


class SyntheticRows:
    """Generate one Parquet batch at a time without holding 500k dicts in RAM."""

    def __init__(self, count: int):
        self.count = count

    def __len__(self) -> int:
        return self.count

    def __getitem__(self, item):
        if isinstance(item, slice):
            start, stop, step = item.indices(self.count)
            return [self[index] for index in range(start, stop, step)]
        if item < 0:
            item += self.count
        if not 0 <= item < self.count:
            raise IndexError(item)
        amount_1 = round(100 + item % 100_000 / 100, 2)
        amount_2 = round(amount_1 + 1.25, 2)
        return {
            "Primary Key": f"INV-{item:07d}",
            "Field": "AMOUNT",
            "File 1 Value": amount_1,
            "File 2 Value": amount_2,
            "Difference": round(amount_2 - amount_1, 2),
            "File Pair ID": "benchmark-pair",
            "Sheet Rule ID": "benchmark-rule",
        }


def configure_isolated_runtime(root: Path) -> None:
    """Set safe test paths and prevent Pydantic from reading any .env file."""
    root.mkdir(parents=True, exist_ok=True)
    os.environ.update({
        "DATABASE_URL": f"sqlite:///{(root / 'acceptance.db').as_posix()}",
        "LOCAL_STORAGE_PATH": str(root / "storage"),
        "STORAGE_BACKEND": "local",
        "JOB_QUEUE_BACKEND": "local",
        "ENVIRONMENT": "test",
        "LOG_LEVEL": "WARNING",
        "LOG_FORMAT": "text",
        "SENTRY_DSN": "",
        "BOOTSTRAP_ADMIN_USERNAME": "",
        "BOOTSTRAP_ADMIN_EMAIL": "",
        "BOOTSTRAP_ADMIN_PASSWORD": "",
        "BOOTSTRAP_ADMIN_TOTP_SECRET": "",
        "SMTP_HOST": "",
        "SMTP_USERNAME": "",
        "SMTP_PASSWORD": "",
        "SMTP_FROM_EMAIL": "",
    })
    from app.core.config import Settings, get_settings

    # Startup diagnostics inspect this path, so use a nonexistent file instead
    # of None. This keeps the real backend/.env completely out of the test.
    Settings.model_config["env_file"] = root / ".phase2-acceptance-env-disabled"
    get_settings.cache_clear()


def _rss_mb(process) -> float:
    return process.memory_info().rss / 1_000_000


def _percentile95(values: list[float]) -> float:
    ordered = sorted(values)
    return ordered[max(0, math.ceil(len(ordered) * 0.95) - 1)]


def run_results_api_check(temp_root: Path, row_count: int, repetitions: int) -> dict[str, Any]:
    try:
        import psutil
        from fastapi.testclient import TestClient
    except ImportError as exc:
        raise SystemExit("Missing API benchmark dependencies. Activate backend/.venv and install backend/requirements.txt.") from exc

    configure_isolated_runtime(temp_root)
    from app.api.routes.reports import _job_report
    from app.database.base import Base
    from app.database.session import SessionLocal, engine
    from app.main import app
    from app.models.job import ReconciliationJob
    from app.models.report import Report
    from app.services.report_data_store import write_report_tables
    from app.storage.base import report_data_directory

    Base.metadata.create_all(bind=engine)
    session_id = str(uuid4())
    report_storage_path = f"reports/{session_id}/phase2-acceptance.xlsx"
    storage_root = Path(os.environ["LOCAL_STORAGE_PATH"])
    data_directory = storage_root / report_data_directory(report_storage_path)
    write_report_tables({
        "metadata": {"file_1_name": "Benchmark source.xlsx", "file_2_name": "Benchmark destination.xlsx"},
        "statistics": {"source_records": row_count, "report_rows": row_count},
        "overall_status": "EXCEPTIONS FOUND",
        "control_checks": [],
        "sheet_rules": [],
        "file_pair_reports": {},
        "exceptions": SyntheticRows(row_count),
        "matched_records": [],
        "missing_in_file_1": [],
        "missing_in_file_2": [],
        "field_differences": [],
        "identity_resolution": [],
        "auto_resolved": [],
    }, data_directory)

    report_summary = {
        "source_records": row_count,
        "source_records_file_1": row_count,
        "source_records_file_2": row_count,
        "report_rows": row_count,
        "matched_records": 0,
        "fully_matched_records": 0,
        "only_in_file_1": 0,
        "only_in_file_2": 0,
        "confidence_review": 0,
    }
    with SessionLocal() as db:
        report = Report(
            session_id=session_id,
            filename="phase2-acceptance.xlsx",
            storage_backend="local",
            storage_path=report_storage_path,
            summary_json=json.dumps(report_summary),
        )
        db.add(report)
        db.flush()
        job = ReconciliationJob(
            session_id=session_id,
            job_type="generic",
            display_name="Phase 2 acceptance fixture",
            status="completed",
            progress=100,
            report_id=report.id,
        )
        db.add(job)
        db.commit()
        job_id = job.id

    process = psutil.Process()
    headers = {"X-Session-ID": session_id}
    summary_url = f"/api/reports/job/{job_id}/summary"
    preview_url = f"/api/reports/job/{job_id}/preview"

    with TestClient(app) as client:
        warm_summary = client.get(summary_url, headers=headers)
        warm_preview = client.get(preview_url, params={"category": "discrepancies"}, headers=headers)
        if warm_summary.status_code != 200 or warm_preview.status_code != 200:
            raise RuntimeError(
                f"API fixture failed: summary={warm_summary.status_code}, preview={warm_preview.status_code}; "
                f"{warm_summary.text[:300]} {warm_preview.text[:300]}"
            )
        if warm_preview.json().get("total_rows") != row_count or len(warm_preview.json().get("rows", [])) != 25:
            raise RuntimeError("Preview returned unexpected row totals or page size.")

        gc.collect()
        rss_before = _rss_mb(process)
        samples = [rss_before]
        sample_done = threading.Event()

        def sample_memory() -> None:
            while not sample_done.wait(0.01):
                samples.append(_rss_mb(process))

        sampler = threading.Thread(target=sample_memory, daemon=True)
        sampler.start()
        timings_ms: list[float] = []
        for _ in range(repetitions):
            started = time.perf_counter()
            summary_response = client.get(summary_url, headers=headers)
            preview_response = client.get(preview_url, params={"category": "discrepancies"}, headers=headers)
            elapsed_ms = (time.perf_counter() - started) * 1000
            if summary_response.status_code != 200 or preview_response.status_code != 200:
                raise RuntimeError("A repeated summary/preview request failed.")
            timings_ms.append(elapsed_ms)
            del summary_response, preview_response
        sample_done.set()
        sampler.join(timeout=2)
        gc.collect()
        rss_after = _rss_mb(process)
        samples.append(rss_after)

    result = {
        "check": "results_api_500k",
        "rows_in_parquet": row_count,
        "page_size": 25,
        "repetitions": repetitions,
        "summary_plus_preview_median_ms": round(statistics.median(timings_ms), 1),
        "summary_plus_preview_p95_ms": round(_percentile95(timings_ms), 1),
        "summary_plus_preview_max_ms": round(max(timings_ms), 1),
        "under_one_second_p95": _percentile95(timings_ms) < 1000,
        "rss_before_mb": round(rss_before, 1),
        "rss_after_mb": round(rss_after, 1),
        "rss_delta_mb": round(rss_after - rss_before, 1),
        "peak_sampled_rss_mb": round(max(samples), 1),
        "peak_sampled_delta_mb": round(max(samples) - rss_before, 1),
        "memory_growth_threshold": "not specified in the plan; review the measured deltas",
        "method": "FastAPI TestClient in an isolated process with a temporary SQLite DB and Parquet result",
    }
    print(json.dumps(result), flush=True)
    return result


def _normal_cell(value: Any) -> Any:
    if isinstance(value, str):
        return re.sub(r"Generated \d{1,2} [A-Za-z]{3} \d{4} \d{2}:\d{2}", "Generated <timestamp>", value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def _table_data(workbook, sheet_name: str) -> tuple[tuple[Any, ...], list[tuple[Any, ...]]]:
    sheet = workbook[sheet_name]
    headers = tuple(_normal_cell(value) for value in next(
        sheet.iter_rows(min_row=4, max_row=4, values_only=True)
    ))
    while headers and headers[-1] is None:
        headers = headers[:-1]
    rows = []
    if headers:
        for row in sheet.iter_rows(min_row=5, max_col=len(headers), values_only=True):
            values = tuple(_normal_cell(value) for value in row)
            if any(value is not None for value in values):
                rows.append(values)
    return headers, rows


def compare_workbooks(previous: Path, current: Path) -> dict[str, Any]:
    """Compare reconciliation data independently of workbook presentation."""
    import openpyxl

    old_book = openpyxl.load_workbook(previous, read_only=True, data_only=True)
    new_book = openpyxl.load_workbook(current, read_only=True, data_only=True)
    differences: list[dict[str, Any]] = []
    compared: dict[str, dict[str, int]] = {}
    detail_sheets = ("02 Differences", "03 Only in File 1", "04 Only in File 2", "06 Matched")
    old_sheets, new_sheets = set(old_book.sheetnames), set(new_book.sheetnames)
    for sheet_name in detail_sheets:
        if sheet_name not in old_sheets or sheet_name not in new_sheets:
            differences.append({"sheet": sheet_name, "issue": "detail sheet missing from one workbook"})
            continue
        old_headers, old_rows = _table_data(old_book, sheet_name)
        new_headers, new_rows = _table_data(new_book, sheet_name)
        compared[sheet_name] = {"previous_rows": len(old_rows), "current_rows": len(new_rows)}
        if old_headers != new_headers:
            differences.append({"sheet": sheet_name, "issue": "detail table headers differ"})
        if old_rows != new_rows:
            first_difference = next((
                index for index, pair in enumerate(zip(old_rows, new_rows), start=1) if pair[0] != pair[1]
            ), min(len(old_rows), len(new_rows)) + 1)
            differences.append({
                "sheet": sheet_name,
                "issue": "detail table values differ",
                "first_differing_data_row": first_difference,
                "previous_row_count": len(old_rows),
                "current_row_count": len(new_rows),
                "previous_sample": str(old_rows[first_difference - 1])[:240] if first_difference <= len(old_rows) else None,
                "current_sample": str(new_rows[first_difference - 1])[:240] if first_difference <= len(new_rows) else None,
            })

    check_name = "07 Checks"
    core_checks = {
        "Records compared",
        "Matched – all fields agree",
        "Matched – values differ",
        "Not found in the other file",
    }
    if check_name not in old_sheets or check_name not in new_sheets:
        differences.append({"sheet": check_name, "issue": "control totals sheet missing from one workbook"})
    else:
        def checks(workbook) -> dict[str, tuple[Any, ...]]:
            sheet = workbook[check_name]
            return {
                str(row[0]): tuple(_normal_cell(value) for value in row[1:3])
                for row in sheet.iter_rows(min_row=5, values_only=True)
                if row and isinstance(row[0], str)
            }

        old_checks, new_checks = checks(old_book), checks(new_book)
        compared_checks = sorted(core_checks & old_checks.keys() & new_checks.keys())
        if compared_checks != sorted(core_checks):
            differences.append({
                "sheet": check_name,
                "issue": "one or more required control totals are missing",
                "missing_previous": sorted(core_checks - old_checks.keys()),
                "missing_current": sorted(core_checks - new_checks.keys()),
            })
        for label in compared_checks:
            if old_checks[label] != new_checks[label]:
                differences.append({
                    "sheet": check_name,
                    "check": label,
                    "previous_counts": old_checks[label],
                    "current_counts": new_checks[label],
                })
    old_book.close()
    new_book.close()
    return {
        "check": "report_reconciliation_data_parity",
        "previous_writer": "openpyxl at commit 146e82a^",
        "current_writer": "xlsxwriter at current checkout",
        "compared_detail_sheets": compared,
        "control_totals_compared": sorted(core_checks),
        "report_presentation": "The earlier summary and detail-sheet layout is restored; filterable styled ranges replace native Excel tables for constant-memory writing.",
        "reconciliation_data_identical": not differences,
        "differences": differences,
        "manual_visual_review_required": True,
    }


def run_report_parity_check(row_count: int, output_directory: Path) -> dict[str, Any]:
    import bench

    from app.reconciliation_engine.engine import run_generic_reconciliation
    from app.reconciliation_engine.preprocessing.vector_preprocessing import prepare_dataframe
    from app.reconciliation_engine.universal_reporter import generate_enterprise_report

    output_directory.mkdir(parents=True, exist_ok=True)
    books, bank = (prepare_dataframe(frame) for frame in bench.generate(row_count))
    result = run_generic_reconciliation(
        books.copy(), bank.copy(), output_directory / "unused.xlsx",
        write_report=False, **bench.SCENARIOS["keyed"],
    )
    universal_data = result["universal_data"]
    current_path = output_directory / "current_writer.xlsx"
    previous_path = output_directory / "previous_writer.xlsx"
    generate_enterprise_report(universal_data, {}, current_path)

    legacy_source = subprocess.run(
        ["git", "show", "146e82a^:backend/app/reconciliation_engine/universal_reporter.py"],
        cwd=REPO,
        check=True,
        capture_output=True,
    ).stdout.decode("utf-8")
    with tempfile.TemporaryDirectory(prefix="recliq-legacy-writer-") as legacy_temp:
        legacy_module_path = Path(legacy_temp) / "legacy_universal_reporter.py"
        legacy_module_path.write_text(legacy_source, encoding="utf-8")
        spec = importlib.util.spec_from_file_location("recliq_legacy_universal_reporter", legacy_module_path)
        if spec is None or spec.loader is None:
            raise RuntimeError("Could not load the pre-Phase-2 report writer from Git history.")
        legacy_module = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = legacy_module
        spec.loader.exec_module(legacy_module)
        legacy_module.generate_enterprise_report(universal_data, {}, previous_path)

    parity = compare_workbooks(previous_path, current_path)
    parity["benchmark_rows_per_input"] = row_count
    parity["artifact_directory"] = str(output_directory)
    parity["review_files"] = [str(previous_path), str(current_path)]
    print(json.dumps(parity), flush=True)
    return parity


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--api-rows", type=int, default=500_000)
    parser.add_argument("--report-rows", type=int, default=4_000)
    parser.add_argument("--api-repetitions", type=int, default=20)
    args = parser.parse_args()
    if args.api_rows < 25 or args.report_rows < 100 or args.api_repetitions < 5:
        parser.error("Use at least 25 API rows, 100 report rows, and 5 API repetitions.")

    with tempfile.TemporaryDirectory(prefix="recliq-phase2-api-") as temp_name:
        try:
            api_result = run_results_api_check(Path(temp_name), args.api_rows, args.api_repetitions)
        finally:
            database_module = sys.modules.get("app.database.session")
            if database_module is not None:
                database_module.engine.dispose()
    report_result = run_report_parity_check(args.report_rows, PERF / "out" / "phase2-acceptance")
    summary = {
        "api_latency_gate": "PASS" if api_result["under_one_second_p95"] else "FAIL",
        "api_memory": {
            "rss_delta_mb": api_result["rss_delta_mb"],
            "peak_sampled_delta_mb": api_result["peak_sampled_delta_mb"],
            "status": "review measured growth; the plan does not set a numeric threshold",
        },
        "report_data_parity": "PASS" if report_result["reconciliation_data_identical"] else "FAIL",
        "report_visual_approval": "OWNER REVIEW REQUIRED",
        "report_artifacts": report_result["artifact_directory"],
    }
    print(json.dumps({"acceptance_summary": summary}, indent=2), flush=True)
    if not api_result["under_one_second_p95"] or not report_result["reconciliation_data_identical"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
