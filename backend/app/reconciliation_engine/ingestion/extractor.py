import hashlib
import json
import os
import tempfile
import threading
from collections import OrderedDict
from concurrent.futures import Future, ProcessPoolExecutor
from pathlib import Path
from typing import List, Dict, Any, Optional
import pandas as pd

try:
    import docx
except ImportError:
    docx = None


def get_file_extension(filename: str) -> str:
    return os.path.splitext(filename)[1].lower()


_format_cache: "OrderedDict[tuple, str | None]" = OrderedDict()


def finance_format(file_path: Path, filename: str) -> str | None:
    """MT940/CAMT.053/BAI2/GSTR-2B, or None for an ordinary table.

    ``.txt`` files are sniffed, since banks deliver MT940 and BAI2 as .txt.
    """
    from app.reconciliation_engine.ingestion import finance_formats

    ext = get_file_extension(filename)
    if ext not in finance_formats.EXTENSION_FORMATS and ext != ".txt":
        return None
    key = _cache_key(Path(file_path), "__format__")
    if key is not None and key in _format_cache:
        return _format_cache[key]
    try:
        with open(file_path, "rb") as handle:
            head = handle.read(8192)
    except OSError:
        return None
    fmt = finance_formats.format_for(ext, head)
    if key is not None:
        _format_cache[key] = fmt
        while len(_format_cache) > 256:
            _format_cache.popitem(last=False)
    return fmt


# .xls (Excel 97-2003) needs xlrd. Calamine is faster for .xlsx, with openpyxl
# retained as a compatibility fallback for workbooks it cannot read.
_EXCEL_ENGINES = {".xlsx": "calamine", ".xls": "xlrd"}


def _excel_engines(extension: str) -> tuple[str, ...]:
    engine = _EXCEL_ENGINES[extension]
    return (engine, "openpyxl") if extension == ".xlsx" else (engine,)


def _read_excel_table(file_path: Path, extension: str, **kwargs) -> pd.DataFrame:
    last_error = None
    for engine in _excel_engines(extension):
        try:
            return pd.read_excel(file_path, engine=engine, **kwargs)
        except Exception as exc:
            last_error = exc
    if last_error is not None:
        raise last_error
    raise ValueError(f"No Excel reader is configured for {extension} files.")


def _read_excel_sheet_names(file_path: Path, extension: str) -> list[str]:
    last_error = None
    for engine in _excel_engines(extension):
        try:
            with pd.ExcelFile(file_path, engine=engine) as workbook:
                return workbook.sheet_names
        except Exception as exc:
            last_error = exc
    if last_error is not None:
        raise last_error
    raise ValueError(f"No Excel reader is configured for {extension} files.")


def _max_rows() -> int | None:
    try:
        from app.core.config import get_settings

        return get_settings().max_rows_per_sheet
    except Exception:
        return None


def _enforce_row_limit(df: pd.DataFrame, sheet_id: str) -> pd.DataFrame:
    """Backstop for the upload-time check (e.g. PDF/Word tables)."""
    limit = _max_rows()
    if limit and len(df) > limit:
        raise ValueError(f"Sheet '{sheet_id}' has {len(df):,} rows; the limit is {limit:,} rows per sheet.")
    return df


def extract_file_metadata(file_path: Path, filename: str) -> List[Dict[str, Any]]:
    """
    Returns a list of available 'sheets' or 'tables' within the file.
    For CSV, it's just one table.
    For Excel, it's the sheet names.
    For PDF/Word, it attempts to find tables.
    """
    ext = get_file_extension(filename)

    if fmt := finance_format(file_path, filename):
        from app.reconciliation_engine.ingestion.finance_formats import sheet_names

        return [{"id": name, "name": name} for name in sheet_names(fmt)]

    if ext in _EXCEL_ENGINES:
        try:
            return [{"id": sheet, "name": sheet} for sheet in _read_excel_sheet_names(file_path, ext)]
        except Exception as e:
            raise ValueError(f"Failed to read Excel file: {e}")
            
    elif ext in ['.csv', '.tsv', '.txt']:
        # For CSV/TSV, there is only one "sheet".
        return [{"id": "default", "name": "Main Data"}]
        
    elif ext == '.pdf':
        # One entry per logical table: a table that runs over several pages is
        # merged into one (see pdf_tables), not listed once per page.
        from app.reconciliation_engine.ingestion.pdf_tables import extract_tables

        return [{"id": table.id, "name": table.name} for table in extract_tables(file_path)]
        
    elif ext == '.docx':
        if docx is None:
            raise ImportError("python-docx is required to parse Word documents.")
        try:
            doc = docx.Document(file_path)
            tables = []
            for i, table in enumerate(doc.tables):
                # Basic check if it looks like a data table (has rows)
                if len(table.rows) > 0:
                    tables.append({
                        "id": f"table_{i+1}",
                        "name": f"Table {i+1} ({len(table.rows)} rows)"
                    })
            if not tables:
                raise ValueError("No data tables detected in Word document.")
            return tables
        except Exception as e:
            raise ValueError(f"Failed to parse Word document: {e}")
            
    else:
        raise ValueError(f"Unsupported file format: {ext}")


# ── Parsed-sheet cache ─────────────────────────────────────────────────────
# Parsed sheets are cached beside uploads as column-readable Parquet. The
# compact facts file lets setup screens inspect headers without loading rows.
_CACHEABLE_EXTENSIONS = {".xlsx", ".xls", ".csv", ".txt", ".tsv", ".sta", ".mt940", ".940", ".bai", ".bai2", ".xml", ".json", ".pdf"}
_cache_guard = threading.Lock()
_parse_locks: dict[tuple, threading.Lock] = {}
# Bump whenever parsing or the cache format changes.
_READER_VERSION = 5


def _cache_path(file_path: Path, sheet_id: str) -> Path:
    digest = hashlib.sha1(str(sheet_id).encode("utf-8")).hexdigest()[:12]
    return file_path.with_name(f"{file_path.name}.sheet-{digest}-r{_READER_VERSION}.parquet")


def _facts_path(parquet_path: Path) -> Path:
    return parquet_path.with_suffix(".facts.json")


def _cache_key(file_path: Path, sheet_id: str) -> tuple | None:
    try:
        stat = file_path.stat()
    except OSError:
        return None
    return (str(file_path.resolve()), stat.st_mtime_ns, stat.st_size, str(sheet_id))


def _is_null_scalar(value: Any) -> bool:
    if value is None:
        return True
    try:
        result = pd.isna(value)
        return isinstance(result, (bool, type(pd.NA))) and bool(result)
    except (TypeError, ValueError):
        return False


def _column_type(series: pd.Series) -> str:
    kinds: set[str] = set()
    for value in series:
        if _is_null_scalar(value):
            continue
        if isinstance(value, (bool,)):
            kinds.add("boolean")
        elif isinstance(value, (pd.Timestamp,)) or hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
            kinds.add("date")
        elif pd.api.types.is_number(value):
            kinds.add("numeric")
        else:
            kinds.add("text")
        if len(kinds) > 1:
            return "mixed"
    return next(iter(kinds), "empty")


def _write_parquet(df: pd.DataFrame, disk_path: Path) -> None:
    """Atomically publish a typed Parquet cache and its non-executable facts."""
    partial = disk_path.with_name(f"{disk_path.name}.{os.getpid()}.part")
    facts_path = _facts_path(disk_path)
    facts_partial = facts_path.with_name(f"{facts_path.name}.{os.getpid()}.part")
    stored = df.copy()
    column_facts = []
    for column in stored.columns:
        kind = _column_type(stored[column])
        if kind == "mixed":
            stored[column] = stored[column].map(lambda value: None if _is_null_scalar(value) else str(value))
        column_facts.append({"name": str(column), "type": kind})
    facts = {"row_count": len(stored), "columns": column_facts, "reader_version": _READER_VERSION}
    try:
        stored.to_parquet(partial, index=False, compression="zstd", engine="pyarrow")
        with facts_partial.open("w", encoding="utf-8") as handle:
            json.dump(facts, handle, ensure_ascii=False, separators=(",", ":"))
        os.replace(partial, disk_path)
        os.replace(facts_partial, facts_path)
    finally:
        partial.unlink(missing_ok=True)
        facts_partial.unlink(missing_ok=True)


def _read_facts(disk_path: Path) -> dict[str, Any] | None:
    try:
        with _facts_path(disk_path).open(encoding="utf-8") as handle:
            facts = json.load(handle)
        return facts if facts.get("reader_version") == _READER_VERSION else None
    except (OSError, ValueError, AttributeError):
        return None


def _read_parquet(disk_path: Path, columns: list[str] | None = None) -> pd.DataFrame:
    available = None
    facts = _read_facts(disk_path)
    if facts is not None:
        available = {item["name"] for item in facts.get("columns", [])}
    if available is None:
        import pyarrow.parquet as pq

        available = set(pq.ParquetFile(disk_path).schema.names)
    selected = [column for column in dict.fromkeys(columns or []) if column in available] if columns is not None else None
    return pd.read_parquet(disk_path, columns=selected, engine="pyarrow")


def read_table_data(file_path: Path, filename: str, sheet_id: str, columns: list[str] | None = None) -> pd.DataFrame:
    """Read one sheet/table, parsing each workbook sheet at most once.

    ``columns`` projects a rule's working set at the Parquet reader. No full
    DataFrame is retained in the API process between requests.
    """
    key = _cache_key(file_path, sheet_id) if get_file_extension(filename) in _CACHEABLE_EXTENSIONS else None
    if key is None:
        frame = _parse_table_data(file_path, filename, sheet_id)
        return frame if columns is None else frame[[column for column in columns if column in frame.columns]]

    with _cache_guard:
        lock = _parse_locks.setdefault(key, threading.Lock())

    # Concurrent requests wait for the first parse instead of repeating it.
    with lock:
        disk_path = _cache_path(file_path, sheet_id)
        if not disk_path.exists():
            _await_warm_up(file_path)
        if disk_path.exists() and disk_path.stat().st_mtime_ns >= key[1]:
            try:
                frame = _read_parquet(disk_path, columns)
                return _enforce_row_limit(frame, sheet_id)
            except Exception:
                disk_path.unlink(missing_ok=True)
                _facts_path(disk_path).unlink(missing_ok=True)
        frame = _parse_table_data(file_path, filename, sheet_id)
        try:
            _write_parquet(frame, disk_path)
        except OSError:
            pass  # Read-only or temporary storage: parse in memory for this call.
        return frame if columns is None else frame[[column for column in columns if column in frame.columns]]


def read_table_sample(file_path: Path, filename: str, sheet_id: str, nrows: int) -> pd.DataFrame:
    """First ``nrows`` rows of a sheet, without waiting for a full parse.

    Setup analysis only needs a sample; reading it directly takes about a
    second even while the full sheet is still being parsed in the background.
    """
    disk_path = _cache_path(file_path, sheet_id)
    key = _cache_key(file_path, sheet_id)
    if key and disk_path.exists() and disk_path.stat().st_mtime_ns >= key[1]:
        try:
            import duckdb

            with duckdb.connect(config={"threads": "1"}) as connection:
                return connection.execute(
                    "SELECT * FROM read_parquet(?) LIMIT ?", [str(disk_path), max(0, nrows)]
                ).df()
        except Exception:
            pass
    ext = get_file_extension(filename)
    if ext in _EXCEL_ENGINES:
        return _read_excel_table(file_path, ext, sheet_name=sheet_id, nrows=nrows)
    if ext in [".csv", ".txt", ".tsv"] and not finance_format(file_path, filename):
        return pd.read_csv(file_path, sep="\t" if ext == ".tsv" else ",", nrows=nrows, dtype=object)
    return read_table_data(file_path, filename, sheet_id).head(nrows)


def read_table_columns(file_path: Path, filename: str, sheet_id: str) -> list[str]:
    """Return a sheet's header row without parsing its data rows when possible."""
    key = _cache_key(file_path, sheet_id)
    disk_path = _cache_path(file_path, sheet_id)
    if key and disk_path.exists() and disk_path.stat().st_mtime_ns >= key[1]:
        facts = _read_facts(disk_path)
        if facts is not None:
            return [item["name"] for item in facts.get("columns", [])]
        try:
            import pyarrow.parquet as pq

            return [str(column) for column in pq.ParquetFile(disk_path).schema.names]
        except Exception:
            pass
    ext = get_file_extension(filename)
    if ext in _EXCEL_ENGINES:
        return [str(column) for column in _read_excel_table(file_path, ext, sheet_name=sheet_id, nrows=0).columns]
    if ext in [".csv", ".txt", ".tsv"] and not finance_format(file_path, filename):
        return [str(column) for column in pd.read_csv(file_path, sep="\t" if ext == ".tsv" else ",", nrows=0).columns]
    return [str(column) for column in read_table_data(file_path, filename, sheet_id).columns]


def _parse_sheets_to_disk(file_path: str, filename: str) -> None:
    """Worker-process entry point: write every sheet's parsed Parquet cache."""
    path = Path(file_path)
    for sheet in extract_file_metadata(path, filename):
        disk_path = _cache_path(path, sheet["id"])
        if not disk_path.exists():
            _write_parquet(_parse_table_data(path, filename, sheet["id"]), disk_path)


# Parsing is CPU-bound. In a thread it would hold the GIL and slow every
# interactive request (sheet lists, headers, analysis) while it runs, so
# warm-up runs in a separate process that only writes the disk cache.
_warm_pool: ProcessPoolExecutor | None = None
_warm_jobs: dict[str, Future] = {}


def warm_table_cache(file_path: Path, filename: str) -> None:
    """Parse every sheet of an upload in a background process, ahead of use."""
    global _warm_pool
    try:
        with _cache_guard:
            if _warm_pool is None:
                _warm_pool = ProcessPoolExecutor(max_workers=2)
            _warm_jobs[str(file_path.resolve())] = _warm_pool.submit(_parse_sheets_to_disk, str(file_path), filename)
    except Exception:
        pass  # Warm-up is best effort; reads parse on demand.


def _await_warm_up(file_path: Path, timeout: float = 120) -> None:
    """If a warm-up is parsing this file, wait for it instead of parsing twice."""
    job = _warm_jobs.get(str(file_path.resolve()))
    if job is not None:
        try:
            job.result(timeout=timeout)
        except Exception:
            pass


def clear_table_cache(file_path: Path) -> None:
    """Drop on-disk parsed sheets when their upload is deleted."""
    for cached in file_path.parent.glob(f"{file_path.name}.sheet-*"):
        if cached.suffix in {".pkl", ".parquet", ".json", ".part"}:
            cached.unlink(missing_ok=True)


def _parse_table_data(file_path: Path, filename: str, sheet_id: str) -> pd.DataFrame:
    """Read one sheet/table and apply the configured row limit."""
    return _enforce_row_limit(_parse_table_unchecked(file_path, filename, sheet_id), sheet_id)


def _parse_table_unchecked(file_path: Path, filename: str, sheet_id: str) -> pd.DataFrame:
    ext = get_file_extension(filename)

    if fmt := finance_format(file_path, filename):
        from app.reconciliation_engine.ingestion.finance_formats import parse_file

        sheets = parse_file(file_path, fmt)
        if sheet_id not in sheets:
            raise ValueError(f"Sheet '{sheet_id}' does not exist in this file. Available: {', '.join(sheets)}.")
        return sheets[sheet_id]

    if ext in _EXCEL_ENGINES:
        return _read_excel_table(file_path, ext, sheet_name=sheet_id)
        
    elif ext in ['.csv', '.txt']:
        return pd.read_csv(file_path, dtype=object)
        
    elif ext == '.tsv':
        return pd.read_csv(file_path, sep='\t', dtype=object)
        
    elif ext == '.pdf':
        from app.reconciliation_engine.ingestion.pdf_tables import extract_tables, find_table

        return find_table(extract_tables(file_path), sheet_id).to_dataframe()

    elif ext == '.docx':
        if docx is None:
            raise ImportError("python-docx is required")
            
        try:
            parts = sheet_id.split('_')
            table_idx = int(parts[1]) - 1
            
            doc = docx.Document(file_path)
            if table_idx < len(doc.tables):
                table = doc.tables[table_idx]
                
                data = []
                for row in table.rows:
                    data.append([cell.text for cell in row.cells])
                    
                if not data or len(data) < 2:
                    return pd.DataFrame()
                    
                df = pd.DataFrame(data[1:], columns=data[0])
                return df
            else:
                raise ValueError(f"Table index {table_idx} out of range")
        except Exception as e:
            raise ValueError(f"Failed to extract Word table: {e}")
            
    else:
        raise ValueError(f"Unsupported file format: {ext}")
