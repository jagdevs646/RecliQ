import io
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2] / "backend"
sys.path.insert(0, str(ROOT))

import openpyxl
from fastapi.testclient import TestClient

from app.main import app


def _upload(client: TestClient, name: str, content: str) -> str:
    response = client.post(
        "/api/files/upload",
        files={"file": (name, io.BytesIO(content.encode("utf-8")), "text/csv")},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def test_date_amount_workflow_persists_previews_and_explainable_group_report():
    with TestClient(app) as client:
        source_id = _upload(
            client,
            "Cashbook.csv",
            "DATE,AMOUNT,ACCOUNT\n05/04/2021,100,ACCT-1\n05/04/2021,200,ACCT-1\n",
        )
        destination_id = _upload(
            client,
            "Bank.csv",
            "POSTED,VALUE,ACCOUNT\n2021-04-05,300,ACCT-1\n",
        )
        payload = {
            "orientation": "vertical",
            "file_pairs": [{
                "file_pair_id": "cash-v-bank",
                "source_files": [{"file_id": source_id}],
                "destination_files": [{"file_id": destination_id}],
                "sheet_rules": [{
                    "sheet_rule_id": "cash-rule",
                    "matching_strategy": {
                        "method": "date_amount",
                        "primary_key_source": [],
                        "primary_key_destination": [],
                    },
                    "date_amount": {
                        "date_source": "DATE",
                        "date_destination": "POSTED",
                        "date_format_source": "day_first",
                        "date_format_destination": "auto",
                        "amount_mode_source": "single",
                        "amount_mode_destination": "single",
                        "amount_source": "AMOUNT",
                        "amount_destination": "VALUE",
                        "text_source": "ACCOUNT",
                        "text_destination": "ACCOUNT",
                        "verify_transactions": True,
                        "allow_group_matches": True,
                        "max_group_size": 4,
                    },
                    "report_label": "Cashbook to Bank",
                }],
            }],
        }
        started = client.post("/api/reconciliation/generic", json=payload)
        assert started.status_code == 200, started.text
        job_id = started.json()["id"]

        job = started.json()
        for _ in range(160):
            job = client.get(f"/api/jobs/{job_id}").json()
            if job["status"] in {"completed", "completed_with_errors", "failed"}:
                break
            time.sleep(0.1)
        assert job["status"] == "completed", job.get("error_message")

        monthly = client.get(f"/api/reports/job/{job_id}/preview", params={"category": "monthly_summary"})
        daily = client.get(f"/api/reports/job/{job_id}/preview", params={"category": "daily_summary"})
        assert monthly.status_code == 200
        assert daily.status_code == 200
        assert monthly.json()["total_rows"] == 1
        assert daily.json()["total_rows"] == 1

        download = client.get(f"/api/reports/job/{job_id}/download")
        assert download.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(download.content), read_only=True, data_only=True)
        assert "07 Monthly Summary" in workbook.sheetnames
        assert "08 Daily Summary" in workbook.sheetnames
        matched_values = [
            value
            for row in workbook["06 Matched"].iter_rows(values_only=True)
            for value in row
        ]
        assert "File 1 transaction details" in matched_values
        assert "File 2 transaction details" in matched_values
        assert any(isinstance(value, str) and "AMOUNT='100'" in value for value in matched_values)
        assert any(isinstance(value, str) and "AMOUNT='200'" in value for value in matched_values)
        assert any(isinstance(value, str) and "normalized=2021-04-05" in value for value in matched_values)
        workbook.close()
