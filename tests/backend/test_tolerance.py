"""Tolerance bands accept small differences after matching; Round rounds half up."""
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from app.reconciliation_engine.cache import canonical_date_value
from app.reconciliation_engine.engine import run_generic_reconciliation
from app.reconciliation_engine.preprocessing import prepare_dataframe
from app.reconciliation_engine.transformations import round_half_up
from app.reconciliation_engine.universal_reporter import generate_enterprise_report
from app.schemas.reconciliation import ToleranceBand

FIELDS = [
    {"file_1_fields": ["DEBIT"], "file_2_fields": ["DEBIT"]},
    {"file_1_fields": ["VALUE DATE"], "file_2_fields": ["POSTING DATE"]},
    {"file_1_fields": ["NARRATION"], "file_2_fields": ["NARRATION"]},
]


def _bank(reference: str, debit: float, date: str = "10/08/2026", narration: str = "Rent") -> dict:
    return {"UTR": reference, "DEBIT": debit, "VALUE DATE": date, "NARRATION": narration}


def _books(reference: str, debit: float, date: str = "10/08/2026", narration: str = "Rent") -> dict:
    return {"REFERENCE": reference, "DEBIT": debit, "POSTING DATE": date, "NARRATION": narration}


def _reco(tmp_path: Path, bank: list[dict], books: list[dict], tolerances: list[dict], **options) -> dict:
    return run_generic_reconciliation(
        prepare_dataframe(pd.DataFrame(bank)), prepare_dataframe(pd.DataFrame(books)), tmp_path / "out.xlsx",
        ["UTR"], ["REFERENCE"], FIELDS, write_report=False, tolerances=tolerances, **options,
    )


def _matched(result: dict) -> dict:
    return {row["MATCH KEY"]: row for row in result["universal_data"]["matched_records"]}


def _differences(result: dict) -> dict:
    rows: dict = {}
    for row in result["universal_data"]["exceptions"]:
        rows.setdefault(row["Primary Key"], {})[row["Field"]] = row
    return rows


# ── Amount, percentage and days ───────────────────────────────────────────
def test_a_difference_within_the_amount_tolerance_moves_the_record_to_matched_with_a_comment(tmp_path: Path):
    result = _reco(tmp_path, [_bank("U1", 169516.16), _bank("U2", 60588.69)], [_books("U1", 169516.2), _books("U2", 60584.71)],
                   [{"field": "DEBIT", "amount": 0.05}])
    assert _matched(result)["U1"]["WITHIN TOLERANCE"] == "DEBIT differs by 0.04, within ±0.05"
    # 3.98 is beyond the tolerance: left exactly as it was.
    assert _differences(result)["U2"]["DEBIT"]["Difference"] == pytest.approx(3.98)
    summary = result["summary"]
    assert summary["fully_matched_records"] == 1 and summary["field_discrepancies"] == 1
    assert summary["within_tolerance_records"] == 1 and summary["within_tolerance_fields"] == 1


def test_the_exact_limit_is_within_tolerance(tmp_path: Path):
    result = _reco(tmp_path, [_bank("U1", 100.05)], [_books("U1", 100.00)], [{"field": "DEBIT", "amount": 0.05}])
    assert "U1" in _matched(result)


def test_a_percentage_tolerance_is_measured_against_the_other_files_value(tmp_path: Path):
    result = _reco(tmp_path, [_bank("U1", 1010), _bank("U2", 1011)], [_books("U1", 1000), _books("U2", 1000)],
                   [{"field": "DEBIT", "percent": 1}])
    assert _matched(result)["U1"]["WITHIN TOLERANCE"] == "DEBIT differs by 10.00 (1.00%), within ±1%"
    assert "DEBIT" in _differences(result)["U2"]


def test_amount_or_percentage_either_limit_is_enough(tmp_path: Path):
    result = _reco(tmp_path, [_bank("U1", 5.9), _bank("U2", 100_400)], [_books("U1", 5), _books("U2", 100_000)],
                   [{"field": "DEBIT", "amount": 1, "percent": 0.5}])
    assert {"U1", "U2"} <= set(_matched(result))


def test_a_days_tolerance_accepts_nearby_dates(tmp_path: Path):
    result = _reco(tmp_path,
                   [_bank("U1", 50, "12/08/2026"), _bank("U2", 50, "20/08/2026")],
                   [_books("U1", 50, "10/08/2026"), _books("U2", 50, "10/08/2026")],
                   [{"field": "VALUE DATE", "days": 3}])
    assert _matched(result)["U1"]["WITHIN TOLERANCE"] == "VALUE DATE differs by 2 days, within ±3 days"
    assert "VALUE DATE" in _differences(result)["U2"]


def test_a_field_can_be_named_by_its_destination_column(tmp_path: Path):
    result = _reco(tmp_path, [_bank("U1", 50, "11/08/2026")], [_books("U1", 50, "10/08/2026")], [{"field": "Posting Date", "days": 1}])
    assert "U1" in _matched(result)


def test_every_field_band_applies_to_numbers_and_dates_but_never_to_text(tmp_path: Path):
    result = _reco(tmp_path,
                   [_bank("U1", 100.02, "11/08/2026"), _bank("U2", 100.02, "11/08/2026", narration="Lease rent")],
                   [_books("U1", 100, "10/08/2026"), _books("U2", 100, "10/08/2026")],
                   [{"field": "*", "amount": 0.05, "days": 1}])
    assert _matched(result)["U1"]["WITHIN TOLERANCE"] == "DEBIT differs by 0.02, within ±0.05; VALUE DATE differs by 1 day, within ±1 day"
    remaining = _differences(result)["U2"]
    assert list(remaining) == ["NARRATION"]  # Only the text difference is left…
    assert remaining["NARRATION"]["Comment"].startswith("Other differences accepted within tolerance: DEBIT differs by 0.02")


def test_a_match_that_needs_confirming_is_not_confirmed_by_a_tolerance(tmp_path: Path):
    # Paired by the date-and-amount pass, not by the key: still for a person to confirm.
    passes = [{"type": "amount_tolerance", "amount_source": "DEBIT", "amount_destination": "DEBIT", "amount_tolerance": 1}]
    result = _reco(tmp_path, [_bank("U1", 100.02)], [_books("X9", 100)], [{"field": "DEBIT", "amount": 0.05}], matching_passes=passes)
    assert result["summary"]["exception_matches"] == 1
    assert not result["universal_data"]["matched_records"]
    assert result["universal_data"]["exceptions"] == []  # The difference itself is accepted.


def test_without_bands_nothing_changes(tmp_path: Path):
    result = _reco(tmp_path, [_bank("U1", 100.02)], [_books("U1", 100)], [])
    assert "DEBIT" in _differences(result)["U1"] and not _matched(result)


def test_a_band_on_a_field_that_is_not_compared_is_rejected(tmp_path: Path):
    with pytest.raises(ValueError, match="not one of the compared fields"):
        _reco(tmp_path, [_bank("U1", 1)], [_books("U1", 1)], [{"field": "CREDIT", "amount": 1}])


def test_a_band_needs_a_limit():
    with pytest.raises(ValueError, match="amount, a percentage, a number of days or a text similarity"):
        ToleranceBand(field="DEBIT")
    assert ToleranceBand(field="DEBIT", percent=0.5).percent == 0.5


# ── Text similarity ────────────────────────────────────────────────────────
def test_text_at_least_as_similar_as_required_is_accepted(tmp_path: Path):
    result = _reco(tmp_path,
                   [_bank("U1", 50, narration="ORBIT TELECOM SETTLEMENT"), _bank("U2", 50, narration="ACME")],
                   [_books("U1", 50, narration="ORBIT TELECOM Telecom Expense"), _books("U2", 50, narration="ZENITH SERVICES")],
                   [{"field": "NARRATION", "similarity": 75}])
    # The same score the report shows in its Similarity column (80%).
    assert _matched(result)["U1"]["WITHIN TOLERANCE"] == "NARRATION is 80% similar, at least 75% required"
    assert "NARRATION" in _differences(result)["U2"]  # 57% similar.


def test_dates_are_never_accepted_for_looking_alike(tmp_path: Path):
    result = _reco(tmp_path, [_bank("U1", 50, "20/08/2026")], [_books("U1", 50, "10/08/2026")], [{"field": "*", "similarity": 10}])
    assert "VALUE DATE" in _differences(result)["U1"]


def test_spreadsheet_dates_read_as_text_keep_their_month(tmp_path: Path):
    # Excel dates reach the comparison as "2026-08-05 00:00:00"; day-first
    # parsing used to turn them into 8 May, 30 days from "2026-08-04".
    result = _reco(tmp_path, [_bank("U1", 50, pd.Timestamp("2026-08-05"))], [_books("U1", 50, pd.Timestamp("2026-08-04"))],
                   [{"field": "VALUE DATE", "days": 3}])
    assert _matched(result)["U1"]["WITHIN TOLERANCE"] == "VALUE DATE differs by 1 day, within ±3 days"


@pytest.mark.parametrize("text", ["2026-08-05 00:00:00", "2026/08/05", "2026.08.05", "2026-08-05T10:30:00", "05/08/2026 10:30"])
def test_year_first_and_timed_dates_parse_the_same_way_whatever_the_convention(text: str):
    expected = date(2026, 8, 5)
    assert canonical_date_value(text, dayfirst=True) == expected
    if text[0:4] == "2026":
        assert canonical_date_value(text, dayfirst=False) == expected


# ── The report ─────────────────────────────────────────────────────────────
def test_the_report_explains_each_accepted_difference(tmp_path: Path):
    result = _reco(tmp_path, [_bank("U1", 100.03), _bank("U2", 150)], [_books("U1", 100), _books("U2", 100)],
                   [{"field": "DEBIT", "amount": 0.05}])
    path = tmp_path / "report.xlsx"
    generate_enterprise_report(result["universal_data"], {}, path)
    workbook = openpyxl.load_workbook(path)
    matched = list(workbook["06 Matched"].iter_rows(min_row=4, values_only=True))
    assert "Comment" in matched[0]
    assert matched[1][matched[0].index("Comment")] == "Matched within tolerance: DEBIT differs by 0.03, within ±0.05"
    summary = [row for row in workbook["01 Summary"].iter_rows(values_only=True) if row[1] == "Tolerance (after matching)"]
    assert summary and "1 difference(s) accepted; 1 record(s) moved" in summary[0][2]


# ── Round ──────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("value, decimals, expected", [
    (168064.55, 1, 168064.6),   # round() gives 168064.5 (the float is 168064.5499…)
    (2.675, 2, 2.68),           # round() gives 2.67
    (0.5, 0, 1.0),              # round() gives 0 (banker's rounding)
    (-2.5, 0, -3.0),            # halves go away from zero
    (169516.16, 1, 169516.2),
])
def test_round_goes_half_up_as_written(value: float, decimals: int, expected: float):
    assert round_half_up(value, decimals) == expected
