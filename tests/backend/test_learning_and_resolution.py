"""Learning from reviewers' decisions, and auto-resolution rules."""
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from _helpers import preview, single_rule_plan, upload, wait_for_job
from app.main import app
from app.reconciliation_engine.engine import run_generic_reconciliation
from app.reconciliation_engine.learning import build_context, wilson_lower_bound
from app.reconciliation_engine.preprocessing import prepare_dataframe
from app.reconciliation_engine.resolution import compile_rule, describe_rule, text_signature
from app.reconciliation_engine.universal_reporter import generate_enterprise_report

VENDOR_KEYED = dict(
    key_file_1=["Vendor"], key_file_2=["Vendor"],
    secondary_conditions=[{"source_column": "Invoice", "destination_column": "Invoice", "comparison_method": "exact_text"}],
)
AMOUNT_DATE_PASS = {"type": "amount_date", "amount_source": "Amount", "amount_destination": "Amount",
                    "date_source": "Date", "date_destination": "Date", "date_window_days": 3}


def _run(tmp_path: Path, source: list[dict], destination: list[dict], **options) -> dict:
    options.setdefault("rules", [{"file_1_fields": ["Amount"], "file_2_fields": ["Amount"]}])
    return run_generic_reconciliation(
        prepare_dataframe(pd.DataFrame(source)), prepare_dataframe(pd.DataFrame(destination)),
        tmp_path / "out.xlsx", write_report=False, **options,
    )


def _decision(value_1, value_2, decision, method=None, confidence=None, minutes=0):
    created = datetime(2026, 9, 1, tzinfo=timezone.utc) + timedelta(minutes=minutes)
    return SimpleNamespace(value_1=value_1, value_2=value_2, decision=decision, method=method, confidence=confidence, created_at=created)


def _identity(result: dict) -> list[dict]:
    return result["universal_data"]["identity_resolution"]


# ── Learning in the engine ─────────────────────────────────────────────────
SIMILAR_SOURCE = [{"Vendor": "Globex Trading Company", "Invoice": "A1", "Amount": 100}]
SIMILAR_DESTINATION = [{"Vendor": "Globex Tradng Company", "Invoice": "A1", "Amount": 100}]


def test_similar_key_match_is_proposed_without_history(tmp_path: Path):
    identity = _identity(_run(tmp_path, SIMILAR_SOURCE, SIMILAR_DESTINATION, **VENDOR_KEYED))[0]
    assert identity["IDENTITY CLASSIFICATION"] == "EXCEPTION_MATCH"
    assert identity["MATCH METHOD"] == "Similar key (company name)"
    assert "REVIEW HISTORY" not in identity


def test_a_rejected_pairing_is_not_proposed_again(tmp_path: Path):
    learning = build_context([_decision("Globex Trading Company", "Globex Tradng Company", "reject")])
    result = _run(tmp_path, SIMILAR_SOURCE, SIMILAR_DESTINATION, learning=learning, **VENDOR_KEYED)
    identity = _identity(result)[0]
    assert identity["IDENTITY CLASSIFICATION"] == "NOT_FOUND"
    assert "rejected pairing this record with 'Globex Tradng Company' before on 2026-09-01" in identity["MATCH EXPLANATION"]
    assert result["summary"]["only_in_file_1"] == 1 and result["summary"]["only_in_file_2"] == 1


def test_a_reset_withdraws_the_rejection(tmp_path: Path):
    learning = build_context([
        _decision("Globex Trading Company", "Globex Tradng Company", "reject"),
        _decision("Globex Trading Company", "Globex Tradng Company", "reset", minutes=5),
    ])
    identity = _identity(_run(tmp_path, SIMILAR_SOURCE, SIMILAR_DESTINATION, learning=learning, **VENDOR_KEYED))[0]
    assert identity["IDENTITY CLASSIFICATION"] == "EXCEPTION_MATCH"


def test_rejecting_one_of_two_candidates_lets_the_other_be_proposed(tmp_path: Path):
    destination = [
        {"Vendor": "Globex Tradng Company", "Invoice": "A1", "Amount": 100},
        {"Vendor": "Globex Trading Compny", "Invoice": "A1", "Amount": 100},
    ]
    without = _identity(_run(tmp_path, SIMILAR_SOURCE, destination, **VENDOR_KEYED))[0]
    assert without["IDENTITY CLASSIFICATION"] == "AMBIGUOUS_MATCH"  # Never picks between two.
    learning = build_context([_decision("Globex Trading Company", "Globex Tradng Company", "reject")])
    identity = _identity(_run(tmp_path, SIMILAR_SOURCE, destination, learning=learning, **VENDOR_KEYED))[0]
    assert identity["IDENTITY CLASSIFICATION"] == "EXCEPTION_MATCH"  # Still for a person to confirm.
    assert identity["CANDIDATE KEY"] == "Globex Trading Compny"
    assert "rejected pairing" in identity["MATCH EXPLANATION"]


def test_review_history_and_learned_confidence_are_explained(tmp_path: Path):
    method = "Similar key (company name)"
    decisions = [_decision("Globex Trading Company", "Globex Tradng Company", "accept", method, 96, minutes=0)]
    decisions += [_decision(f"Vendor {i}", f"Vendr {i}", "accept", method, 96, minutes=i) for i in range(1, 9)]
    decisions += [_decision("Other", "Othr", "reject", method, 97, minutes=20)]
    confidence = int(str(_identity(_run(tmp_path, SIMILAR_SOURCE, SIMILAR_DESTINATION, **VENDOR_KEYED))[0]["MATCH CONFIDENCE"]).rstrip("%"))
    assert 95 <= confidence <= 99, confidence  # Same 95-99% band as the decisions above.
    learning = build_context(decisions)
    identity = _identity(_run(tmp_path, SIMILAR_SOURCE, SIMILAR_DESTINATION, learning=learning, **VENDOR_KEYED))[0]
    assert identity["REVIEW HISTORY"] == "Confirmed by a reviewer 1 time before"
    assert identity["LEARNED CONFIDENCE"] == "90% (9 of 10 confirmed)"
    assert identity["IDENTITY CLASSIFICATION"] == "EXCEPTION_MATCH"  # Learning never confirms by itself.


def test_rejected_keyless_pairing_is_skipped(tmp_path: Path):
    source = [{"Ref": "R2", "Date": "05/09/2026", "Amount": 750}]
    destination = [{"Ref": "BANK-77", "Date": "07/09/2026", "Amount": 750}]
    options = dict(key_file_1=["Ref"], key_file_2=["Ref"], matching_passes=[AMOUNT_DATE_PASS])
    assert _run(tmp_path, source, destination, **options)["summary"]["keyless_matches"] == 1
    learning = build_context([_decision("R2", "BANK-77", "reject", "Pass 2: Amount + date ±3 days")])
    result = _run(tmp_path, source, destination, learning=learning, **options)
    assert result["summary"]["keyless_matches"] == 0
    assert "Pass 2: Amount + date ±3 days: a reviewer rejected" in _identity(result)[0]["MATCH EXPLANATION"]


def test_wilson_lower_bound_is_conservative():
    assert wilson_lower_bound(10, 10) < 1.0
    assert wilson_lower_bound(48, 50) > 0.85 > wilson_lower_bound(5, 5)


# ── Rule validation ────────────────────────────────────────────────────────
BANK_CHARGES = {
    "name": "Bank charges", "category": "only_in_destination",
    "conditions": [
        {"column": "Narrative", "operator": "contains", "value": "bank charges"},
        {"column": "Amount", "operator": "abs_lte", "value": 50},
    ],
    "action": {"resolution": "Bank charges", "reason_code": "BCHG", "gl_account": "6100"},
}


@pytest.mark.parametrize(
    "change, message",
    [
        ({"conditions": []}, "at least one condition"),
        ({"category": "field_difference", "conditions": [{"operator": "field_is", "value": "AMOUNT"}]}, "limit how large"),
        ({"category": "to_confirm", "conditions": [{"operator": "method_is", "value": "Pass 2"}]}, "earlier reviewer confirmation"),
        ({"conditions": [{"column": "*", "operator": "not_contains", "value": "x"}]}, "'Any column' only works"),
        ({"conditions": [{"column": "Amount", "operator": "between", "value": 10, "value_2": 5}]}, "must not be smaller"),
        ({"conditions": [{"operator": "difference_abs_lte", "value": 1}]}, "cannot be used"),
        ({"action": {"resolution": ""}}, "Say how matching exceptions are resolved"),
    ],
)
def test_unsafe_or_incomplete_rules_are_refused(change, message):
    with pytest.raises(ValueError, match=message):
        compile_rule({**BANK_CHARGES, **change})


def test_rule_reads_as_one_sentence():
    text = describe_rule(compile_rule(BANK_CHARGES))
    assert text == ("Only in the destination file where Narrative contains 'bank charges' and Amount is at most "
                    "(ignoring sign) 50 → Bank charges.")


def test_text_signature_ignores_numbers_and_identifiers():
    assert text_signature("BANK CHARGES JAN-2026 #4411") == "bank charges jan"
    assert text_signature("INV-000123") is None
    assert text_signature("1500.00") is None


# ── Auto-resolution in the engine ──────────────────────────────────────────
BOOKS = [
    {"Ref": "INV-1", "Amount": 100.00, "Narrative": "Customer payment"},
    {"Ref": "INV-2", "Amount": 250.00, "Narrative": "Customer payment"},
]
BANK = [
    {"Ref": "INV-1", "Amount": 100.40, "Narrative": "Customer payment"},
    {"Ref": "INV-2", "Amount": 250.00, "Narrative": "Customer payment"},
    {"Ref": "CHG-1", "Amount": -12.50, "Narrative": "BANK CHARGES JAN"},
    {"Ref": "CHG-2", "Amount": -95.00, "Narrative": "Bank charges wire"},
    {"Ref": "UNK-1", "Amount": -30.00, "Narrative": "Unknown debit"},
]
ROUNDING = {
    "name": "Rounding", "category": "field_difference",
    "conditions": [{"operator": "field_is", "value": "AMOUNT"}, {"operator": "difference_abs_lte", "value": 0.5}],
    "action": {"resolution": "Rounding difference", "gl_account": "6999"},
}


def _resolved_run(tmp_path: Path, rules: list[dict]) -> dict:
    return _run(
        tmp_path, BOOKS, BANK, key_file_1=["Ref"], key_file_2=["Ref"],
        include_columns_file_2=["Narrative"], resolution_rules=rules,
    )


def test_rules_resolve_recurring_exceptions_with_an_audit_trail(tmp_path: Path):
    result = _resolved_run(tmp_path, [BANK_CHARGES, ROUNDING])
    summary, data = result["summary"], result["universal_data"]
    assert summary["only_in_file_2"] == 2  # The 95.00 charge is above the cap; the unknown debit stays.
    assert summary["auto_resolved_only_in_file_2"] == 1 and summary["auto_resolved_differences"] == 1
    assert summary["field_discrepancies"] == 0  # INV-1's 0.40 difference is explained.
    assert summary["fully_matched_records"] == 2
    charge = next(row for row in data["auto_resolved"] if row["__KIND__"] == "only_in_destination")
    assert charge["Rule"] == "Bank charges" and charge["Rule Version"] == 1 and charge["GL Account"] == "6100"
    assert charge["Record"] == "CHG-1" and charge["NARRATIVE"] == "BANK CHARGES JAN"
    rounding = next(row for row in data["auto_resolved"] if row["__KIND__"] == "field_difference")
    assert rounding["Field"] == "AMOUNT" and rounding["Difference"] == pytest.approx(-0.4)
    matched = next(row for row in data["matched_records"] if row.get("MATCH KEY") == "INV-1")
    assert "AMOUNT: Rounding difference" in matched["RESOLVED DIFFERENCES"]
    assert data["metadata"]["resolution_rules_used"] == {"Bank charges (v1)": 1, "Rounding (v1)": 1}


def test_first_rule_in_priority_order_wins_and_disabled_rules_do_nothing(tmp_path: Path):
    catch_all = {**BANK_CHARGES, "name": "Any debit", "conditions": [{"column": "Amount", "operator": "lte", "value": 0}],
                 "action": {"resolution": "Other debit"}}
    result = _resolved_run(tmp_path, [BANK_CHARGES, catch_all])
    rules = [row["Rule"] for row in result["universal_data"]["auto_resolved"]]
    assert rules == ["Bank charges", "Any debit", "Any debit"]
    assert _resolved_run(tmp_path, [{**BANK_CHARGES, "enabled": False}])["summary"]["auto_resolved"] == 0


def test_a_condition_on_a_missing_column_never_resolves(tmp_path: Path):
    rule = {**BANK_CHARGES, "conditions": [{"column": "Memo", "operator": "not_contains", "value": "x"}]}
    assert _resolved_run(tmp_path, [rule])["summary"]["auto_resolved"] == 0


def test_ambiguous_records_are_never_auto_resolved(tmp_path: Path):
    rule = {"name": "Small items", "category": "only_in_source",
            "conditions": [{"column": "Amount", "operator": "abs_lte", "value": 1000}], "action": {"resolution": "Small"}}
    result = _run(
        tmp_path,
        [{"Date": "01/09/2026", "Amount": 100}, {"Date": "02/09/2026", "Amount": 100}],
        [{"Date": "01/09/2026", "Amount": 100}],
        key_file_1=[], key_file_2=[], matching_passes=[AMOUNT_DATE_PASS], resolution_rules=[rule],
    )
    assert result["summary"]["auto_resolved"] == 0 and result["summary"]["ambiguous_matches"] == 2


def test_confirmation_rule_needs_earlier_reviewer_confirmation(tmp_path: Path):
    rule = {"name": "Reviewed before", "category": "to_confirm",
            "conditions": [{"operator": "previously_confirmed", "value": 2}], "action": {"resolution": "Confirmed by rule"}}
    once = build_context([_decision("Globex Trading Company", "Globex Tradng Company", "accept")])
    result = _run(tmp_path, SIMILAR_SOURCE, SIMILAR_DESTINATION, learning=once, resolution_rules=[rule], **VENDOR_KEYED)
    assert result["summary"]["exception_matches"] == 1 and result["summary"]["auto_confirmed_matches"] == 0

    twice = build_context([_decision("Globex Trading Company", "Globex Tradng Company", "accept", minutes=m) for m in (0, 1)])
    result = _run(tmp_path, SIMILAR_SOURCE, SIMILAR_DESTINATION, learning=twice, resolution_rules=[rule], **VENDOR_KEYED)
    identity = _identity(result)[0]
    assert result["summary"]["auto_confirmed_matches"] == 1 and result["summary"]["exception_matches"] == 0
    assert identity["AUTO-RESOLVED BY"] == "Reviewed before (v1)"
    assert "Confirmed automatically by rule 'Reviewed before (v1)'" in identity["MATCH EXPLANATION"]
    assert result["summary"]["fully_matched_records"] == 1

    rejected = build_context([*[_decision("Globex Trading Company", "Globex Tradng Company", "accept", minutes=m) for m in (0, 1)],
                              _decision("Globex Trading Company", "Globex Tradng Company", "reject", minutes=2)])
    result = _run(tmp_path, SIMILAR_SOURCE, SIMILAR_DESTINATION, learning=rejected, resolution_rules=[rule], **VENDOR_KEYED)
    assert result["summary"]["auto_confirmed_matches"] == 0  # A later rejection wins.


def test_report_lists_auto_resolved_items_and_checks_still_balance(tmp_path: Path):
    result = _resolved_run(tmp_path, [BANK_CHARGES, ROUNDING])
    path = tmp_path / "report.xlsx"
    generate_enterprise_report(result["universal_data"], {}, path)
    workbook = openpyxl.load_workbook(path)
    assert "09 Auto-resolved" in workbook.sheetnames
    sheet = workbook["09 Auto-resolved"]
    headers = [cell.value for cell in sheet[4]]
    assert headers[:7] == ["Exception", "Resolution", "Rule", "Rule Version", "Reason Code", "GL Account", "Record"]
    checks = {row[0]: row[1:] for row in workbook["07 Checks"].iter_rows(min_row=5, values_only=True) if row[0]}
    assert checks["Not found, explained by an auto-resolution rule"][:2] == (0, 1)
    accounted = next(value for label, value in checks.items() if label.startswith("Records accounted for"))
    assert accounted[2] == "Pass"
    summary_text = " ".join(str(cell.value) for row in workbook["01 Summary"].iter_rows() for cell in row if cell.value)
    assert "Auto-resolved by rules" in summary_text and "Bank charges (v1) (1)" in summary_text


# ── API: rules, audit, preview, decisions, learning, suggestions ──────────
def _bank_job(client: TestClient, charge_narrative: str = "BANK CHARGES JAN", charge: float = -12.5) -> dict:
    books = upload(client, "books.xlsx", {"B": [{"Ref": "INV-1", "Amount": 100}]})
    bank = upload(client, "bank.xlsx", {"S": [
        {"Ref": "INV-1", "Amount": 100, "Narrative": "Payment"},
        {"Ref": "CHG-1", "Amount": charge, "Narrative": charge_narrative},
    ]})
    plan = single_rule_plan(books, bank, "B", "S", **{
        "matching_strategy": {"primary_key_source": ["Ref"], "primary_key_destination": ["Ref"]},
        "reconciliation_mapping": [{"file_1_fields": ["Amount"], "file_2_fields": ["Amount"]}],
        "include_columns_file_2": ["Narrative"],
    })
    return wait_for_job(client, client.post("/api/reconciliation/generic", json=plan).json()["id"])


def test_rule_changes_are_versioned_and_audited_with_before_and_after():
    with TestClient(app) as client:
        created = client.post("/api/resolution-rules", json=BANK_CHARGES)
        assert created.status_code == 201, created.text
        rule = created.json()
        assert rule["version"] == 1 and rule["summary"].startswith("Only in the destination file where Narrative contains")

        refused = client.put(f"/api/resolution-rules/{rule['id']}", json={"conditions": []})
        assert refused.status_code == 422
        updated = client.put(f"/api/resolution-rules/{rule['id']}", json={"action": {**BANK_CHARGES["action"], "gl_account": "6150"}}).json()
        assert updated["version"] == 2 and updated["action"]["gl_account"] == "6150"
        same = client.put(f"/api/resolution-rules/{rule['id']}", json={"action": {**BANK_CHARGES["action"], "gl_account": "6150"}}).json()
        assert same["version"] == 2  # No change, no new version.
        archived = client.delete(f"/api/resolution-rules/{rule['id']}").json()
        assert archived["archived"] is True and client.get("/api/resolution-rules").json()["rules"] == []

        events = client.get("/api/audit/events", params={"entity_id": rule["id"]}).json()["events"]
        actions = [event["action"] for event in events]
        assert {"rule.created", "rule.updated", "rule.archived"} <= set(actions)
        change = next(event for event in events if event["action"] == "rule.updated")
        assert change["before"]["action"]["gl_account"] == "6100" and change["after"]["action"]["gl_account"] == "6150"
        assert client.get("/api/audit/verify").json()["valid"] is True


def test_run_applies_rules_audits_them_and_preview_shows_the_effect():
    with TestClient(app) as client:
        before = _bank_job(client)
        assert client.get(f"/api/reports/job/{before['id']}/summary").json()["only_in_file_2"] == 1
        dry_run = client.post("/api/resolution-rules/preview", json={"job_id": before["id"], "rule": BANK_CHARGES}).json()
        assert dry_run["matches"] == 1 and dry_run["sample"][0]["REF"] == "CHG-1"

        client.post("/api/resolution-rules", json=BANK_CHARGES)
        job = _bank_job(client)
        summary = client.get(f"/api/reports/job/{job['id']}/summary").json()
        assert summary["only_in_file_2"] == 0 and summary["auto_resolved"] == 1
        assert summary["resolution_rules"][0]["name"] == "Bank charges"
        rows = preview(client, job["id"], "auto_resolved")
        assert rows[0]["Resolution"] == "Bank charges" and "__KIND__" not in rows[0]
        events = client.get("/api/audit/events", params={"action": "exceptions.auto_resolved"}).json()["events"]
        assert events and events[0]["entity_id"] == job["id"]
        assert events[0]["after"]["rules"]["Bank charges (v1)"]["count"] == 1


def test_recurring_exceptions_become_rule_suggestions_until_a_rule_covers_them():
    with TestClient(app) as client:
        for _ in range(3):
            _bank_job(client, "SMS ALERT CHARGES QTR", -5.9)
        suggestions = client.get("/api/resolution-rules/suggestions").json()["suggestions"]
        recurring = [item for item in suggestions if item["kind"] == "recurring_exception"]
        assert recurring and recurring[0]["rule"]["conditions"][0]["value"] == "sms alert charges qtr"
        cap = recurring[0]["rule"]["conditions"][1]
        assert cap["operator"] == "abs_lte" and cap["value"] == 6.0  # Capped at the largest amount seen.
        assert recurring[0]["evidence"]["jobs"] == 3

        created = client.post("/api/resolution-rules", json={**recurring[0]["rule"], "action": {"resolution": "SMS charges"}})
        assert created.status_code == 201, created.text
        suggestions = client.get("/api/resolution-rules/suggestions").json()["suggestions"]
        assert not [item for item in suggestions if item["kind"] == "recurring_exception"]


def test_decisions_teach_weights_and_keyless_pairs_never_become_aliases():
    with TestClient(app) as client:
        method = "Pass 2: Amount + date ±3 days"
        # 40 unanimous confirmations: enough evidence (95% lower bound > 90%).
        for index in range(40):
            response = client.post("/api/aliases/decisions", json={
                "value_1": f"R{index}", "value_2": f"BANK-{index}", "decision": "accept",
                "job_id": f"job-{index % 3}", "method": method, "confidence": 99,
            })
            assert response.status_code == 201
        client.post("/api/aliases/decisions", json={"value_1": "R1", "value_2": "BANK-1", "decision": "accept", "job_id": "job-9", "method": method})
        learning = client.get("/api/learning").json()
        weight = next(row for row in learning["method_weights"] if row["method"] == method)
        assert weight["accepted"] == 41 and weight["band"] == "all" and weight["enough_evidence"]
        kinds = {item["kind"] for item in learning["suggestions"]}
        assert "auto_confirm_method" in kinds
        # R1 ↔ BANK-1 was confirmed in two runs, but amount/date pairs are not names.
        assert client.get("/api/aliases/suggestions").json()["suggestions"] == []

        client.post("/api/aliases/decisions", json={"value_1": "R1", "value_2": "BANK-1", "decision": "reject", "method": method})
        remembered = client.get("/api/learning").json()["remembered_rejections"]
        assert [(row["value_1"], row["value_2"]) for row in remembered] == [("R1", "BANK-1")]
        reset = client.post("/api/aliases/decisions", json={"value_1": "R1", "value_2": "BANK-1", "decision": "reset"})
        assert reset.status_code == 201
        assert client.get("/api/learning").json()["remembered_rejections"] == []
        actions = {event["action"] for event in client.get("/api/audit/events").json()["events"]}
        assert {"match.accepted", "match.rejected", "match.reset"} <= actions


def test_rule_catalog_lists_allowed_conditions_per_exception_type():
    with TestClient(app) as client:
        catalog = client.get("/api/resolution-rules/catalog").json()
    categories = {item["id"] for item in catalog["categories"]}
    assert categories == {"only_in_source", "only_in_destination", "field_difference", "to_confirm"}
    difference = next(item for item in catalog["operators"] if item["id"] == "difference_abs_lte")
    assert difference["categories"] == ["field_difference"]
