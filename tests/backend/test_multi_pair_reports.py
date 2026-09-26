import io
import sys
import time
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2] / "backend"
sys.path.insert(0, str(ROOT))

import openpyxl
import pandas as pd
from fastapi.testclient import TestClient

from app.database.session import SessionLocal
from app.main import app
from app.models.history import ReconciliationHistory


def _workbook_bytes(sheets: dict[str, list[dict]]) -> io.BytesIO:
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        for sheet_name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=sheet_name, index=False)
    stream.seek(0)
    return stream


def test_multiple_file_pairs_download_as_zip_with_independent_workbooks():
    with TestClient(app) as client:
        source = _workbook_bytes({
            "North": [{"Invoice": "N-001", "Amount": 10}],
            "South": [{"Reference": "S-001", "Amount": 20}],
        })
        destination = _workbook_bytes({
            "North ledger": [{"Document": "N001", "Net": 10}],
            "South ledger": [{"Bank Ref": "S001", "Credit": 20}],
        })
        uploaded_source = client.post("/api/files/upload", files={"file": ("North_South.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        uploaded_destination = client.post("/api/files/upload", files={"file": ("Ledger.xlsx", destination, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert uploaded_source.status_code == 201
        assert uploaded_destination.status_code == 201

        source_id = uploaded_source.json()["id"]
        destination_id = uploaded_destination.json()["id"]
        payload = {
            "orientation": "vertical",
            "file_pairs": [
                {
                    "file_pair_id": "north",
                    "source_files": [{"file_id": source_id, "sheet_id": "North"}],
                    "destination_files": [{"file_id": destination_id, "sheet_id": "North ledger"}],
                    "sheet_rules": [{
                        "sheet_rule_id": "north-rule",
                        "source_sheets": ["North"],
                        "destination_sheets": ["North ledger"],
                        "matching_strategy": {"primary_key_source": ["Invoice"], "primary_key_destination": ["Document"]},
                        "reconciliation_mapping": [{"file_1_fields": ["Amount"], "file_2_fields": ["Net"]}],
                        "report_label": "North report",
                    }],
                },
                {
                    "file_pair_id": "south",
                    "source_files": [{"file_id": source_id, "sheet_id": "South"}],
                    "destination_files": [{"file_id": destination_id, "sheet_id": "South ledger"}],
                    "sheet_rules": [{
                        "sheet_rule_id": "south-rule",
                        "source_sheets": ["South"],
                        "destination_sheets": ["South ledger"],
                        "matching_strategy": {"primary_key_source": ["Reference"], "primary_key_destination": ["Bank Ref"]},
                        "reconciliation_mapping": [{"file_1_fields": ["Amount"], "file_2_fields": ["Credit"]}],
                        "report_label": "South report",
                    }],
                },
            ],
        }
        started = client.post("/api/reconciliation/generic", json=payload)
        assert started.status_code == 200
        job = started.json()
        for _ in range(40):
            job = client.get(f"/api/jobs/{job['id']}").json()
            if job["status"] in {"completed", "failed"}:
                break
            time.sleep(0.1)
        assert job["status"] == "completed", job.get("error_message")

        report = client.get(f"/api/reports/job/{job['id']}/download")
        assert report.status_code == 200
        assert report.headers["content-type"].startswith("application/zip")
        assert len(report.content) > 100
        with zipfile.ZipFile(io.BytesIO(report.content)) as archive:
            assert archive.namelist() == ["North_report.xlsx", "South_report.xlsx"]


def test_one_failed_sheet_rule_keeps_successful_results_available():
    with TestClient(app) as client:
        source = _workbook_bytes({
            "Good": [{"Invoice": "G-001", "Amount": 10}],
            "Bad": [{"Reference": "B-001", "Amount": 20}],
        })
        destination = _workbook_bytes({
            "Good ledger": [{"Document": "G001", "Net": 10}],
            "Bad ledger": [{"Bank Ref": "B001", "Credit": 20}],
        })
        source_response = client.post("/api/files/upload", files={"file": ("Source.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        destination_response = client.post("/api/files/upload", files={"file": ("Destination.xlsx", destination, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        source_id = source_response.json()["id"]
        destination_id = destination_response.json()["id"]
        pairs = [
            ("good", "Good", "Good ledger", "Invoice", "Document", "Amount", "Net"),
            ("bad", "Bad", "Bad ledger", "Missing Key", "Bank Ref", "Amount", "Credit"),
        ]
        payload = {
            "file_pairs": [{
                "file_pair_id": pair_id,
                "source_files": [{"file_id": source_id, "sheet_id": source_sheet}],
                "destination_files": [{"file_id": destination_id, "sheet_id": destination_sheet}],
                "sheet_rules": [{
                    "sheet_rule_id": f"{pair_id}-rule",
                    "source_sheets": [source_sheet],
                    "destination_sheets": [destination_sheet],
                    "matching_strategy": {"primary_key_source": [source_key], "primary_key_destination": [destination_key]},
                    "reconciliation_mapping": [{"file_1_fields": [source_amount], "file_2_fields": [destination_amount]}],
                    "report_label": pair_id,
                }],
            } for pair_id, source_sheet, destination_sheet, source_key, destination_key, source_amount, destination_amount in pairs],
        }
        started = client.post("/api/reconciliation/generic", json=payload)
        assert started.status_code == 200
        job = started.json()
        for _ in range(40):
            job = client.get(f"/api/jobs/{job['id']}").json()
            if job["status"] in {"completed", "completed_with_errors", "failed"}:
                break
            time.sleep(0.1)

        assert job["status"] == "completed_with_errors"
        assert "missing columns MISSING KEY" in (job["error_message"] or "")
        preview = client.get(f"/api/reports/job/{job['id']}/preview?category=review")
        assert preview.status_code == 200


_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client: TestClient, name: str, sheets: dict[str, list[dict]]) -> str:
    response = client.post("/api/files/upload", files={"file": (name, _workbook_bytes(sheets), _XLSX)})
    assert response.status_code == 201
    return response.json()["id"]


def _wait_for_job(client: TestClient, job_id: str) -> dict:
    job: dict = {}
    for _ in range(80):
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in {"completed", "completed_with_errors", "failed"}:
            break
        time.sleep(0.1)
    return job


def test_two_sheet_rules_keep_independent_keys_conditions_and_counts():
    """Direct regression for the shared global key/rules bug (spec 1.4)."""
    with TestClient(app) as client:
        source_id = _upload(client, "January.xlsx", {
            "Sales": [
                {"Customer": "Alpha Consulting", "Posted": "2026-09-07", "Amount": 100},
                {"Customer": "Beta Traders", "Posted": "2026-09-08", "Amount": 200},
            ],
            "Payments": [
                {"Reference": "P-1", "Entity": "IN", "Paid": 50},
                {"Reference": "P-1", "Entity": "US", "Paid": 70},
            ],
        })
        destination_id = _upload(client, "January_Ledger.xlsx", {
            "Sales Ledger": [
                {"Client": "Alpha Consultng", "Posting Date": "07/09/2026", "Gross": 100},
                {"Client": "Beta Traders", "Posting Date": "08/09/2026", "Gross": 200},
            ],
            "Bank": [
                {"Bank Ref": "P1", "Company": "IN", "Credit": 50},
                {"Bank Ref": "P1", "Company": "US", "Credit": 75},
            ],
        })
        payload = {
            "file_pairs": [{
                "file_pair_id": "january",
                "source_files": [{"file_id": source_id}],
                "destination_files": [{"file_id": destination_id}],
                "sheet_rules": [
                    {
                        "sheet_rule_id": "sales",
                        "source_sheets": ["Sales"],
                        "destination_sheets": ["Sales Ledger"],
                        "matching_strategy": {
                            "primary_key_source": ["Customer"],
                            "primary_key_destination": ["Client"],
                            "secondary_conditions": [
                                {"source_column": "Posted", "destination_column": "Posting Date", "comparison_method": "normalized_date"},
                                {"source_column": "Amount", "destination_column": "Gross", "comparison_method": "numeric_tolerance", "numeric_tolerance": 0},
                            ],
                        },
                        "reconciliation_mapping": [{"file_1_fields": ["Amount"], "file_2_fields": ["Gross"]}],
                        "report_label": "Sales -> Sales Ledger",
                    },
                    {
                        "sheet_rule_id": "payments",
                        "source_sheets": ["Payments"],
                        "destination_sheets": ["Bank"],
                        "matching_strategy": {
                            "primary_key_source": ["Reference", "Entity"],
                            "primary_key_destination": ["Bank Ref", "Company"],
                        },
                        "reconciliation_mapping": [{"file_1_fields": ["Paid"], "file_2_fields": ["Credit"]}],
                        "report_label": "Payments -> Bank",
                    },
                ],
            }],
        }
        started = client.post("/api/reconciliation/generic", json=payload)
        assert started.status_code == 200
        job = _wait_for_job(client, started.json()["id"])
        assert job["status"] == "completed", job.get("error_message")

        summary = client.get(f"/api/reports/job/{job['id']}/summary").json()
        rules = {rule["sheet_rule_id"]: rule for rule in summary["sheet_rules"]}
        assert rules["sales"]["primary_key_source"] == ["Customer"]
        assert len(rules["sales"]["secondary_conditions"]) == 2
        assert rules["payments"]["primary_key_source"] == ["Reference", "Entity"]
        assert rules["payments"]["secondary_conditions"] == []
        assert rules["sales"]["summary"]["exception_matches"] == 1
        assert rules["sales"]["summary"]["exact_matches"] == 1
        # Composite identity: P-1/US matches only P1/US, and its amount differs.
        assert rules["payments"]["summary"]["exact_matches"] == 2
        assert rules["payments"]["summary"]["field_discrepancies"] == 1
        # Aggregate equals the sum of rules: no double counting.
        assert summary["exact_matches"] == 3
        assert summary["file_pairs"][0]["summary"]["exact_matches"] == 3

        scoped = client.get(f"/api/reports/job/{job['id']}/summary", params={"sheet_rule_id": "payments"}).json()
        assert scoped["exception_matches"] == 0 and scoped["field_discrepancies"] == 1
        assert client.get(f"/api/reports/job/{job['id']}/summary", params={"sheet_rule_id": "nope"}).status_code == 404

        exceptions = client.get(f"/api/reports/job/{job['id']}/preview", params={"category": "exception_matches"}).json()
        assert exceptions["total_rows"] == 1
        assert exceptions["rows"][0]["MATCH KEY"] == "Alpha Consulting"
        assert "Sheet Rule ID" not in exceptions["columns"]
        payments_only = client.get(
            f"/api/reports/job/{job['id']}/preview",
            params={"category": "exception_matches", "sheet_rule_id": "payments"},
        ).json()
        assert payments_only["total_rows"] == 0

        # One file pair keeps the direct .xlsx download, now with a rule breakdown.
        report = client.get(f"/api/reports/job/{job['id']}/download")
        assert report.headers["content-type"].startswith(_XLSX)
        workbook = openpyxl.load_workbook(io.BytesIO(report.content))
        rule_sheet = workbook["08 Sheet Rules"]
        headers = [row[0].value for row in rule_sheet.iter_rows() if str(row[0].value or "").startswith("RULE ")]
        assert headers == ["RULE 1 — Sales ↔ Sales Ledger", "RULE 2 — Payments ↔ Bank"]

        with SessionLocal() as db:
            messages = [item.message for item in db.query(ReconciliationHistory).filter(ReconciliationHistory.job_id == job["id"])]
        assert any("Sales → Sales Ledger (rule 1 of 2)" in message for message in messages)
        assert any("Payments → Bank (rule 2 of 2)" in message for message in messages)


def test_three_file_pairs_zip_with_sanitized_unique_names_and_per_pair_download():
    with TestClient(app) as client:
        pairs = []
        for month in ("Jan", "Feb", "Mar"):
            source_id = _upload(client, f"{month}.xlsx", {"Data": [{"Invoice": f"{month}-1", "Amount": 10}]})
            destination_id = _upload(client, f"{month}_Ledger.xlsx", {"Data": [{"Invoice": f"{month}1", "Amount": 10}]})
            pairs.append((month, source_id, destination_id))
        payload = {
            "file_pairs": [{
                "file_pair_id": month.lower(),
                "source_files": [{"file_id": source_id}],
                "destination_files": [{"file_id": destination_id}],
                # Two pairs share a label containing path characters; the third
                # has no label and falls back to its rule label.
                "report_metadata": {"label": "../Q1/Report" if month != "Mar" else ""},
                "sheet_rules": [{
                    "sheet_rule_id": f"{month.lower()}-rule",
                    "source_sheets": ["Data"],
                    "destination_sheets": ["Data"],
                    "matching_strategy": {"primary_key_source": ["Invoice"], "primary_key_destination": ["Invoice"]},
                    "reconciliation_mapping": [{"file_1_fields": ["Amount"], "file_2_fields": ["Amount"]}],
                    "report_label": f"{month} ledger",
                }],
            } for month, source_id, destination_id in pairs],
        }
        started = client.post("/api/reconciliation/generic", json=payload)
        job = _wait_for_job(client, started.json()["id"])
        assert job["status"] == "completed", job.get("error_message")

        report = client.get(f"/api/reports/job/{job['id']}/download")
        assert report.headers["content-type"].startswith("application/zip")
        assert "RecliQ_Reconciliation_Reports.zip" in report.headers["content-disposition"]
        with zipfile.ZipFile(io.BytesIO(report.content)) as archive:
            names = archive.namelist()
        assert names == ["Q1_Report.xlsx", "Q1_Report_2.xlsx", "Mar_ledger.xlsx"]

        manifest = client.get(f"/api/reports/job/{job['id']}/summary").json()["file_pairs"]
        assert [pair["report_filename"] for pair in manifest] == names
        assert [pair["summary"]["exact_matches"] for pair in manifest] == [1, 1, 1]

        feb = client.get(f"/api/reports/job/{job['id']}/download", params={"file_pair_id": "feb"})
        assert feb.status_code == 200
        assert 'filename="Q1_Report_2.xlsx"' in feb.headers["content-disposition"]
        assert "01 Summary" in openpyxl.load_workbook(io.BytesIO(feb.content)).sheetnames
        assert client.get(f"/api/reports/job/{job['id']}/download", params={"file_pair_id": "nope"}).status_code == 404

        # History shows every pair, and each pair's workbook can be customized.
        assert job["file_pair_count"] == 3
        assert client.post(f"/api/reports/job/{job['id']}/download_custom", json={}).status_code == 422
        custom = client.post(
            f"/api/reports/job/{job['id']}/download_custom",
            params={"file_pair_id": "mar"},
            json={"include_matched": False},
        )
        assert custom.status_code == 200
        assert 'filename="Custom_Mar_ledger.xlsx"' in custom.headers["content-disposition"]
        custom_workbook = openpyxl.load_workbook(io.BytesIO(custom.content))
        assert "06 Matched" not in custom_workbook.sheetnames
        # Mar's own figures (1 record), not the job-wide total of 3.
        assert custom_workbook["01 Summary"]["B7"].value == 1
        assert custom_workbook["01 Summary"]["C11"].value == 1
