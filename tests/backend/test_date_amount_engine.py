import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2] / "backend"
sys.path.insert(0, str(ROOT))

import openpyxl
import pandas as pd
import pytest

from app.reconciliation_engine import date_amount as engine
from app.reconciliation_engine.universal_reporter import generate_enterprise_report
from app.schemas.reconciliation import DateAmountConfig, FilePairConfig, FileSource, ReconciliationPlan, SheetRuleConfig, MatchingStrategy


def config(**overrides):
    result = {
        "date_source": "DATE", "date_destination": "DATE",
        "date_format_source": "auto", "date_format_destination": "auto",
        "amount_mode_source": "single", "amount_mode_destination": "single",
        "amount_source": "AMOUNT", "amount_destination": "AMOUNT",
        "debit_source": None, "credit_source": None,
        "debit_destination": None, "credit_destination": None,
        "positive_is_source": "debit", "positive_is_destination": "debit",
        "reverse_direction_source": False, "reverse_direction_destination": False,
        "text_source": None, "text_destination": None,
        "currency_source": None, "currency_destination": None,
        "amount_tolerance": "0", "date_tolerance_days": 0,
        "verify_transactions": True, "allow_group_matches": True, "max_group_size": 4,
    }
    result.update(overrides)
    return result


def reconcile(source, destination, **settings):
    return engine.reconcile_date_amount(
        pd.DataFrame(source), pd.DataFrame(destination), config(**settings),
        file_1_name="Ledger.xlsx", file_2_name="Statement.xlsx",
    )


def test_date_parser_handles_iso_compact_excel_serial_and_infers_unambiguous_format():
    parsed = engine.parse_date_column(["13/04/2026", "03/04/2026", "20260305", 46000], "auto")
    assert parsed.inferred_format == "day_first"
    assert parsed.values[:3] == (date(2026, 4, 13), date(2026, 4, 3), date(2026, 3, 5))
    assert parsed.values[3] is not None
    assert not parsed.errors


def test_date_parser_handles_month_names_leap_day_and_year_end():
    parsed = engine.parse_date_column(["29-Feb-2024", "31 December 2025", "Apr 01, 2021"], "day_first")
    assert parsed.values == (date(2024, 2, 29), date(2025, 12, 31), date(2021, 4, 1))
    assert not parsed.errors


def test_ambiguous_auto_date_is_not_guessed_but_explicit_format_parses_it():
    ambiguous = engine.parse_date_column(["03/04/2026"], "auto")
    assert ambiguous.values == (None,)
    assert "ambiguous" in ambiguous.errors[0].lower()
    assert engine.parse_date_column(["03/04/2026"], "day_first").values == (date(2026, 4, 3),)
    assert engine.parse_date_column(["03/04/2026"], "month_first").values == (date(2026, 3, 4),)


@pytest.mark.parametrize(("value", "format_name", "expected"), [
    ("₹1,23,456.78", "decimal_point", "123456.78"),
    ("(1,234.50)", "decimal_point", "-1234.50"),
    ("1.234,50", "decimal_comma", "1234.50"),
    ("−25.5", "decimal_point", "-25.5"),
    ("USD 20.00", "decimal_point", "20.00"),
])
def test_financial_parser_uses_exact_decimal_and_supports_common_accounting_formats(value, format_name, expected):
    amount, currency = engine.parse_financial_amount(value, format_name)
    assert str(amount) == expected
    if value.startswith("₹"):
        assert currency == "INR"
    if value.startswith("USD"):
        assert currency == "USD"


def test_amount_comparison_does_not_use_binary_float_equality():
    result = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": ["0.10"]},
        {"DATE": ["2026-04-01"], "AMOUNT": ["0.1000001"]},
    )
    assert result["summary"]["verified_match_groups"] == 0
    assert result["summary"]["only_in_file_1"] == 1
    assert result["summary"]["only_in_file_2"] == 1


def test_exact_one_to_one_matches_and_monthly_daily_summaries_are_year_aware():
    result = reconcile(
        {"DATE": ["2025-01-02", "2026-01-02"], "AMOUNT": [100, 200]},
        {"DATE": ["2025-01-02", "2026-01-02"], "AMOUNT": [100, 200]},
    )
    assert result["summary"]["verified_match_groups"] == 2
    assert result["summary"]["monthly_summary_groups"] == 2
    assert {row["Period"] for row in result["universal_data"]["monthly_summary"]} == {"2025-01", "2026-01"}
    assert result["summary"]["aggregate_balanced_months"] == 2


def test_equal_monthly_totals_do_not_force_different_transactions_into_matches():
    result = reconcile(
        {"DATE": ["2021-04-05", "2021-04-10"], "AMOUNT": [100, 200]},
        {"DATE": ["2021-04-05", "2021-04-10"], "AMOUNT": [150, 150]},
    )
    assert result["summary"]["aggregate_balanced_months"] == 1
    assert result["summary"]["verified_match_groups"] == 0
    assert result["summary"]["only_in_file_1"] == 2
    assert result["summary"]["only_in_file_2"] == 2


def test_daily_totals_do_not_force_different_same_day_transactions_into_matches():
    result = reconcile(
        {"DATE": ["2021-04-05", "2021-04-05"], "AMOUNT": [100, 200]},
        {"DATE": ["2021-04-05", "2021-04-05"], "AMOUNT": [150, 150]},
    )
    assert result["universal_data"]["daily_summary"][0]["Aggregate status"] == "Balanced"
    assert result["summary"]["verified_match_groups"] == 0
    assert result["summary"]["only_in_file_1"] == 2
    assert result["summary"]["only_in_file_2"] == 2


def test_grouped_many_to_one_match_requires_exact_amount_sum_and_same_day():
    result = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [10]},
        {"DATE": ["2026-04-01", "2026-04-01"], "AMOUNT": [4, 6]},
    )
    assert result["summary"]["verified_match_groups"] == 1
    match = result["universal_data"]["matched_records"][0]
    assert match["MATCH TYPE"] == "One-to-many"
    assert match["Source row count"] == 1
    assert match["Comparison row count"] == 2
    assert "AMOUNT=4" in match["File 2 transaction details"]
    assert "AMOUNT=6" in match["File 2 transaction details"]


def test_ambiguous_grouping_is_left_for_review_instead_of_forced():
    result = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [10]},
        {"DATE": ["2026-04-01"] * 4, "AMOUNT": [4, 6, 3, 7]},
    )
    assert result["summary"]["verified_match_groups"] == 0
    assert result["summary"]["confidence_review"] == 5
    assert all("ambiguous" in row["MATCH EXPLANATION"].lower() or "grouping" in row["MATCH EXPLANATION"].lower() for row in result["universal_data"]["identity_resolution"])


def test_duplicate_multisets_pair_only_the_shared_count_and_flag_excess():
    result = reconcile(
        {"DATE": ["2026-04-01", "2026-04-01"], "AMOUNT": [10, 10]},
        {"DATE": ["2026-04-01"], "AMOUNT": [10]},
    )
    assert result["summary"]["verified_match_groups"] == 1
    assert result["summary"]["confidence_review"] == 1
    assert result["summary"]["only_in_file_1"] == 1


def test_unique_amount_tolerance_pair_is_accepted_but_equal_candidates_are_reviewed_both_sides():
    unique = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [100]},
        {"DATE": ["2026-04-01"], "AMOUNT": [100.05]},
        amount_tolerance="0.10",
    )
    assert unique["summary"]["verified_match_groups"] == 1
    assert unique["universal_data"]["matched_records"][0]["MATCH TYPE"] == "Tolerance one-to-one"

    ambiguous = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [100]},
        {"DATE": ["2026-04-01", "2026-04-01"], "AMOUNT": [99.95, 100.05]},
        amount_tolerance="0.10",
    )
    assert ambiguous["summary"]["verified_match_groups"] == 0
    assert ambiguous["summary"]["confidence_review"] == 3


def test_date_tolerance_can_pair_across_month_boundary_but_group_matches_cannot():
    result = reconcile(
        {"DATE": ["2026-01-31"], "AMOUNT": [75]},
        {"DATE": ["2026-02-01"], "AMOUNT": [75]},
        date_tolerance_days=1,
    )
    assert result["summary"]["verified_match_groups"] == 1
    assert result["universal_data"]["matched_records"][0]["Date difference (days)"] == 1


def test_direction_required_text_and_currency_are_hard_match_dimensions():
    direction = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [100]},
        {"DATE": ["2026-04-01"], "AMOUNT": [-100]},
    )
    assert direction["summary"]["verified_match_groups"] == 0

    narrative = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [100], "VENDOR": ["Acme"]},
        {"DATE": ["2026-04-01"], "AMOUNT": [100], "VENDOR": ["Other"]},
        text_source="VENDOR", text_destination="VENDOR",
    )
    assert narrative["summary"]["verified_match_groups"] == 0

    currency = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": ["₹100"]},
        {"DATE": ["2026-04-01"], "AMOUNT": ["$100"]},
    )
    assert currency["summary"]["verified_match_groups"] == 0
    assert {row["Aggregate status"] for row in currency["universal_data"]["monthly_summary"]} == {"Currency mismatch"}


def test_utr_text_key_ignores_formatting_but_report_keeps_the_original_values():
    result = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [100], "UTR": ["UTR2600-0814"]},
        {"DATE": ["2026-04-01"], "AMOUNT": [100], "UTR": ["utr2600/0814"]},
        text_source="UTR", text_destination="UTR",
    )

    assert result["summary"]["verified_match_groups"] == 1
    match = result["universal_data"]["matched_records"][0]
    assert match["Normalized text"] == "utr26000814"
    assert "UTR2600-0814" in match["File 1 transaction details"]
    assert "utr2600/0814" in match["File 2 transaction details"]


def test_date_amount_identifier_normalization_supports_real_bank_column_names():
    result = reconcile(
        {"Books_Date": ["2026-08-09"], "Debit": [14503.01], "Credit": [0], "Reference": ["UTR26000814"]},
        {"Bank_Date": ["2026-08-09"], "Debit": [14503.01], "Credit": [0], "Cheque_UTR": ["UTR2600-0814"]},
        date_source="Books_Date", date_destination="Bank_Date",
        amount_mode_source="debit_credit", amount_mode_destination="debit_credit",
        debit_source="Debit", credit_source="Credit",
        debit_destination="Debit", credit_destination="Credit",
        text_source="Reference", text_destination="Cheque_UTR",
    )

    assert result["summary"]["verified_match_groups"] == 1
    match = result["universal_data"]["matched_records"][0]
    assert match["Normalized text"] == "utr26000814"
    assert "UTR26000814" in match["File 1 transaction details"]
    assert "UTR2600-0814" in match["File 2 transaction details"]


def test_date_amount_name_text_keeps_word_boundaries_and_case_whitespace_normalization():
    result = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [100], "VENDOR NAME": ["  Acme   North  "]},
        {"DATE": ["2026-04-01"], "AMOUNT": [100], "VENDOR NAME": ["north acme"]},
        text_source="VENDOR NAME", text_destination="VENDOR NAME",
    )

    assert result["summary"]["verified_match_groups"] == 1
    assert result["universal_data"]["matched_records"][0]["Normalized text"] == "acme north"


def test_signed_amount_convention_and_reverse_direction_are_applied_per_file():
    direct = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [100]},
        {"DATE": ["2026-04-01"], "AMOUNT": [100]},
        positive_is_source="credit",
        positive_is_destination="debit",
    )
    assert direct["summary"]["verified_match_groups"] == 0

    reversed_source = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [100]},
        {"DATE": ["2026-04-01"], "AMOUNT": [-100]},
        reverse_direction_source=True,
    )
    assert reversed_source["summary"]["verified_match_groups"] == 1


def test_separate_debit_credit_columns_normalize_to_signed_amounts():
    result = reconcile(
        {"DATE": ["2026-04-01"], "DR": [None], "CR": [125]},
        {"DATE": ["2026-04-01"], "AMOUNT": [-125]},
        amount_mode_source="debit_credit", debit_source="DR", credit_source="CR",
    )
    assert result["summary"]["verified_match_groups"] == 1
    invalid = engine.inspect_date_amount_inputs(
        pd.DataFrame({"DATE": ["2026-04-01"], "DR": [20], "CR": [5]}),
        pd.DataFrame({"DATE": ["2026-04-01"], "AMOUNT": [15]}),
        config(amount_mode_source="debit_credit", debit_source="DR", credit_source="CR"),
    )
    assert any(item["category"] == "invalid_amount" and item["side"] == "source" for item in invalid)


def test_aggregate_only_reports_balances_without_claiming_transaction_matches():
    result = reconcile(
        {"DATE": ["2026-04-01", "2026-04-02"], "AMOUNT": [40, 60]},
        {"DATE": ["2026-04-01"], "AMOUNT": [100]},
        verify_transactions=False,
    )
    assert result["summary"]["date_amount_summary_only"] == 1
    assert result["summary"]["verified_match_groups"] == 0
    assert result["universal_data"]["monthly_summary"][0]["Aggregate status"] == "Aggregate Balanced — Transactions Not Verified"
    assert result["universal_data"]["monthly_summary"][0]["Transaction verification"] == "Not performed"
    assert result["universal_data"]["matched_records"] == []


def test_precheck_blocks_ambiguous_auto_dates_and_warns_for_invalid_amounts():
    ambiguous = engine.inspect_date_amount_inputs(
        pd.DataFrame({"DATE": ["03/04/2026"], "AMOUNT": [10]}),
        pd.DataFrame({"DATE": ["2026-04-03"], "AMOUNT": [10]}), config(),
    )
    assert any(item["severity"] == "blocker" and item["category"] == "ambiguous_date_format" for item in ambiguous)

    invalid = engine.inspect_date_amount_inputs(
        pd.DataFrame({"DATE": ["2026-04-03"], "AMOUNT": ["bad"]}),
        pd.DataFrame({"DATE": ["2026-04-03"], "AMOUNT": [10]}), config(),
    )
    assert any(item["severity"] == "warning" and item["category"] == "invalid_amount" for item in invalid)
    assert any(item["severity"] == "warning" and item["category"] == "no_valid_amounts" for item in invalid)

    with pytest.raises(ValueError, match="could mean either day/month or month/day"):
        reconcile(
            {"DATE": ["03/04/2026"], "AMOUNT": [10]},
            {"DATE": ["2026-04-03"], "AMOUNT": [10]},
        )


def test_candidate_search_cap_keeps_all_candidates_for_manual_review(monkeypatch):
    monkeypatch.setattr(engine, "_MAX_TOLERANCE_EDGES", 1)
    result = reconcile(
        {"DATE": ["2026-04-01", "2026-04-01"], "AMOUNT": [100, 100]},
        {"DATE": ["2026-04-01", "2026-04-01"], "AMOUNT": [100.01, 100.02]},
        amount_tolerance="1", allow_group_matches=False,
    )
    assert result["summary"]["verified_match_groups"] == 0
    assert result["summary"]["confidence_review"] == 4
    assert all("safety limit" in row["MATCH EXPLANATION"].lower() for row in result["universal_data"]["identity_resolution"])


def test_group_candidate_cap_leaves_rows_for_review(monkeypatch):
    monkeypatch.setattr(engine, "_MAX_GROUP_CANDIDATES", 1)
    result = reconcile(
        {"DATE": ["2026-04-01"], "AMOUNT": [10]},
        {"DATE": ["2026-04-01"] * 4, "AMOUNT": [4, 6, 3, 7]},
    )
    assert result["summary"]["verified_match_groups"] == 0
    assert result["summary"]["confidence_review"] == 5
    assert any("bounded search limit" in row["Why it was not matched"] for row in result["universal_data"]["missing_in_file_2"])


def test_date_amount_plan_has_a_versionable_canonical_execution_contract():
    settings = DateAmountConfig(date_source="DATE", date_destination="DATE", amount_source="AMOUNT", amount_destination="AMOUNT")
    plan = ReconciliationPlan(file_pairs=[FilePairConfig(
        source_files=[FileSource(file_id="source")],
        destination_files=[FileSource(file_id="destination")],
        sheet_rules=[SheetRuleConfig(matching_strategy=MatchingStrategy(method="date_amount"), date_amount=settings)],
    )])
    item = plan.execution_rules()[0]
    assert item["reconciliation_method"] == "date_amount"
    assert item["date_amount"]["date_source"] == "DATE"
    assert item["key_file_1"] == []


def test_date_amount_report_has_separate_month_and_day_tabs_and_clear_aggregate_only_warning(tmp_path):
    result = reconcile(
        {"DATE": ["2026-04-01", "2026-04-02"], "AMOUNT": [40, 60]},
        {"DATE": ["2026-04-01"], "AMOUNT": [100]},
        verify_transactions=False,
    )
    report = tmp_path / "date-amount.xlsx"
    generate_enterprise_report(result["universal_data"], {}, report)
    workbook = openpyxl.load_workbook(report, read_only=True, data_only=True)
    assert "07 Monthly Summary" in workbook.sheetnames
    assert "08 Daily Summary" in workbook.sheetnames
    summary_values = [cell.value for row in workbook["01 Summary"].iter_rows() for cell in row]
    assert any(isinstance(value, str) and "Transactions were not verified" in value for value in summary_values)
    checks_values = [cell.value for row in workbook["09 Checks"].iter_rows() for cell in row]
    assert "Not performed" in checks_values
    workbook.close()
